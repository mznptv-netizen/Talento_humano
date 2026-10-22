"""Importación revisable de .xlsx y .xlsm. Nunca modifica el archivo fuente."""
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re
from relaciones import REGIONAL_OPTIONS, canonical_assignments, summary, ASSIGNMENT_KEYS
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from talento_humano import FIELDS, KEYS, LABELS, DATES, REQUIRED, REGIONS, validate, normalize, display


ALIASES = {
    'numero': ['NO', 'N°', 'NUMERO DE EMPLEADO', 'NUMERO DE EMPLEADO (NO)', 'NUMERO'],
    'nombre': ['NOMBRE', 'NOMBRES Y APELLIDOS', 'APELLIDOS Y NOMBRES'],
    'apellidos': ['APELLIDOS'], 'nombres': ['NOMBRES'],
    'sede': ['SEDE', 'SEDES'], 'modalidad': ['MODALIDAD', 'MODALIDADES'],
    'horas_totales': ['HORAS TOTALES', 'TOTAL HORAS'],
    'periodo_horas': ['PERIODO HORAS', 'PERIODO'],
    'fecha_fin': ['FECHA DE FINALIZACION DEL CONTRATO', 'FECHA FINALIZACION CONTRATO'],
    'fijo': ['FIJO'], 'celular': ['NUMERO DE CELULAR', 'CELULAR'],
    'afp': ['AFILIACION A FONDO DE PENSIONES (AFP)', 'AFP'],
    'riesgo': ['CLASE RIESGO'],
    'fecha_expedicion': ['FECHA EXPEDICION DEL DCTO', 'FECHA EXPEDICION DEL DOCUMENTO', 'FECHA DE EXPEDICION DEL DOCUMENTO (FECHA EXPEDICION DEL DCTO)'],
    'rh': ['RH'], 'sexo': ['SEXO F-M', 'SEXO'],
    'entidad_financiera': ['EN.FINANCIERA', 'ENTIDAD FINANCIERA (EN.FINANCIERA)'],
    'codigo_banco': ['COD.BANCO', 'CODIGO DEL BANCO (COD.BANCO)'],
    'cuenta': ['N° DE CUENTA', 'Nº DE CUENTA', 'NUMERO DE CUENTA', 'CUENTA'],
    'tipo_cuenta': ['T. CUENTA', 'TIPO DE CUENTA', 'TIPO DE CUENTA (T. CUENTA)'],
    'tipo_contrato': ['T. CONTRATO', 'T. CONTREATO', 'TIPO DE CONTRATO (T. CONTREATO)'],
    'prorroga1': ['PRORROGA1'],
    'estatura': ['ESTATURA'], 'salario': ['SALARIO'],
}


def header_key(text):
    return re.sub(r'[^a-z0-9]', '', normalize(text))


def suggest_mapping(headers):
    mapping = {}
    for key, label, _ in FIELDS:
        aliases = {header_key(x) for x in [label, key.replace('_', ' ')] + ALIASES.get(key, [])}
        matches = [i for i, h in enumerate(headers) if header_key(h) in aliases]
        if len(matches) == 1:
            mapping[key] = matches[0]
    return mapping


def cell_text(cell, key, epoch):
    value = cell.value
    if value is None:
        return ''
    if cell.data_type == 'e':
        raise ValueError(f'{LABELS[key]} contiene un error de Excel: {value}.')
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    if key in DATES and isinstance(value, (float, int)) and not isinstance(value, bool):
        from openpyxl.utils.datetime import from_excel
        converted = from_excel(value, epoch)
        if not isinstance(converted, (datetime, date)):
            raise ValueError(f'{LABELS[key]} no contiene una fecha válida.')
        return converted.strftime('%d/%m/%Y')
    if isinstance(value, bool):
        text = 'Sí' if value else 'No'
    elif isinstance(value, (int, float)):
        text = format(Decimal(str(value)), 'f')
        if '.' in text:
            text = text.rstrip('0').rstrip('.')
        if key in ('numero', 'documento', 'cuenta', 'codigo_banco', 'fijo', 'celular'):
            if not float(value).is_integer():
                raise ValueError(f'{LABELS[key]} tiene decimales; guárdalo como texto en Excel.')
            if abs(value) >= 10**15:
                raise ValueError(f'{LABELS[key]} tiene más de 15 dígitos y es numérico. Verifica el original y guárdalo como texto en Excel para evitar pérdida de precisión.')
            if re.fullmatch(r'0+', cell.number_format or ''):
                text = text.zfill(len(cell.number_format))
    else:
        text = str(value).strip()
    if key in DATES and text:
        for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
            try:
                return datetime.strptime(text, fmt).strftime('%d/%m/%Y')
            except ValueError:
                pass
    if key == 'salario':
        text = re.sub(r'(?i)COP|\$|\s', '', text)
        if re.fullmatch(r'\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?', text):
            text = text.replace('.', '')
    if key == 'sexo':
        text = {'FEMENINO': 'F', 'MASCULINO': 'M'}.get(text.upper(), text.upper())
    if key == 'riesgo':
        text = {'I': '1', 'II': '2', 'III': '3', 'IV': '4', 'V': '5'}.get(text.upper(), text)
    if key == 'rh':
        text = text.upper().replace(' ', '')
    if key == 'regional':
        text = {'bogota': 'Bogotá', 'bogota d.c.': 'Bogotá', 'bogota dc': 'Bogotá',
                'cundinamarca': 'Cundinamarca', 'ambas': 'Ambas'}.get(normalize(text), text)
    return text


class ExcelSource:
    def __init__(self, path):
        from openpyxl import load_workbook
        self.path = Path(path)
        self.values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            self.formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        except Exception:
            self.values.close()
            raise

    @property
    def sheets(self):
        return self.values.sheetnames

    def headers(self, sheet, header_row):
        if header_row < 1:
            raise ValueError('La fila de encabezados debe ser 1 o mayor.')
        rows = self.values[sheet].iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        result = [str(x).strip() if x is not None else '' for x in next(rows, ())]
        if not any(result):
            raise ValueError('La fila seleccionada está vacía. Elige la fila que contiene los encabezados.')
        return result

    def records(self, sheet, header_row, mapping, default_region=None):
        if 'documento' not in mapping or not ('nombre' in mapping or {'apellidos','nombres'} <= set(mapping)):
            raise ValueError('Asigna Documento y Nombre completo, o Documento, Apellidos y Nombres por separado.')
        if 'regional' not in mapping and default_region not in REGIONAL_OPTIONS:
            raise ValueError('Asigna la columna Regional o elige la regional para toda la hoja.')
        if len(set(mapping.values())) != len(mapping):
            raise ValueError('Una columna de Excel no puede asignarse a dos campos. Revisa las listas.')
        assignment_groups, assignment_errors = self.assignment_rows()
        records, errors = [], list(assignment_errors)
        max_col = max(mapping.values()) + 1
        cached = self.values[sheet].iter_rows(min_row=header_row + 1, max_col=max_col)
        formulas = self.formulas[sheet].iter_rows(min_row=header_row + 1, max_col=max_col)
        for row_number, (cells, source) in enumerate(zip(cached, formulas), start=header_row + 1):
            if all(source[index].value is None for index in mapping.values()):
                continue
            try:
                data = {}
                for key, index in mapping.items():
                    cell = cells[index]
                    if source[index].data_type == 'f' and cell.value is None:
                        raise ValueError(f'{LABELS[key]} contiene una fórmula sin resultado guardado. Recalcula y guarda el Excel, o pega su resultado como valor.')
                    data[key] = cell_text(cell, key, self.values.epoch)
                if 'regional' not in mapping:
                    data['regional'] = default_region
                data['_asignaciones'] = assignment_groups.get(data.get('documento',''), [])
                canonical = validate(data)
                data['_asignaciones'] = canonical_assignments(data['_asignaciones'], canonical['regional'], canonical['horas_totales'], canonical['periodo_horas'])
                if data['_asignaciones']:
                    data['sede'] = summary(data['_asignaciones'], 'sede')
                    data['modalidad'] = summary(data['_asignaciones'], 'modalidad')
                records.append((row_number, data))
            except (ValueError, OverflowError) as error:
                errors.append(f'Fila {row_number}: {error}')
        documents = {raw.get('documento','') for _,raw in records}
        unmatched = set(assignment_groups) - documents
        if unmatched:
            errors.append('Hoja Asignaciones: hay documentos sin empleado válido en la hoja elegida: ' + ', '.join(sorted(unmatched)))
        return records, errors

    def assignment_rows(self):
        groups, errors = {}, []
        sheet = next((name for name in self.sheets if normalize(name.strip()) == 'asignaciones'), None)
        if sheet is None:
            return groups, errors
        headers = [header_key(v) for v in self.headers(sheet, 1)]
        required = ('documento', *ASSIGNMENT_KEYS)
        if any(headers.count(key) != 1 for key in required):
            return {}, ['Hoja Asignaciones: la fila 1 debe contener Documento, Regional, Sede, Modalidad, Perfil y Horas, sin repetir columnas.']
        indexes = {key:headers.index(key) for key in required}
        cached = self.values[sheet].iter_rows(min_row=2)
        formulas = self.formulas[sheet].iter_rows(min_row=2)
        for number, (cells, source) in enumerate(zip(cached, formulas), 2):
            if all(source[i].value is None for i in indexes.values()):
                continue
            try:
                row = {}
                for key,index in indexes.items():
                    if source[index].data_type == 'f' and cells[index].value is None:
                        raise ValueError('Fórmula sin resultado guardado: ' + key)
                    row[key] = cell_text(cells[index], key, self.values.epoch)
                if not row['documento']:
                    raise ValueError('Falta Documento.')
                groups.setdefault(row.pop('documento'), []).append(row)
            except ValueError as error:
                errors.append(f'Asignaciones, fila {number}: {error}')
        return groups, errors

    def close(self):
        self.values.close()
        self.formulas.close()


class ImportDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app, self.source = app, None
        self.title('Importar empleados desde Excel')
        self.geometry('940x730')
        self.minsize(860, 650)
        self.transient(app)
        self.grab_set()
        self.protocol('WM_DELETE_WINDOW', self.close)
        self.mapping_vars, self.column_indexes = {}, {}
        self.loaded_config = None
        ttk.Label(self, text='Importar empleados', style='Title.TLabel').pack(anchor='w', padx=22, pady=(18, 6))
        ttk.Label(self, text='Selecciona el archivo, revisa las columnas y comprueba los datos antes de guardarlos.').pack(anchor='w', padx=22)
        top = ttk.Frame(self, padding=(22, 12))
        top.pack(fill='x')
        ttk.Button(top, text='1. Seleccionar Excel', command=self.choose_file).grid(row=0, column=0, sticky='w')
        self.filename = tk.StringVar(value='Archivos .xlsx o .xlsm')
        ttk.Label(top, textvariable=self.filename, wraplength=600).grid(row=0, column=1, columnspan=4, sticky='w', padx=12)
        ttk.Label(top, text='Hoja').grid(row=1, column=0, sticky='w', pady=10)
        self.sheet = tk.StringVar()
        self.sheets = ttk.Combobox(top, textvariable=self.sheet, state='readonly', width=25)
        self.sheets.grid(row=1, column=1, sticky='w', padx=12)
        ttk.Label(top, text='Fila de encabezados').grid(row=1, column=2)
        self.header_row = tk.StringVar(value='1')
        ttk.Spinbox(top, from_=1, to=1000, width=5, textvariable=self.header_row).grid(row=1, column=3, padx=10)
        ttk.Button(top, text='2. Leer columnas', command=self.read_columns).grid(row=1, column=4)
        ttk.Label(top, text='Sin columna Regional, usar:').grid(row=2, column=0, columnspan=2, sticky='w', pady=(4, 8))
        self.default_region = tk.StringVar(value=app.active_region if app.active_region in REGIONS else '')
        ttk.Combobox(top, textvariable=self.default_region, values=REGIONAL_OPTIONS, state='readonly', width=20).grid(row=2, column=2, columnspan=2, sticky='w', pady=(4, 8))
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill='both', expand=True, padx=22)
        map_tab = ttk.Frame(self.notebook)
        self.notebook.add(map_tab, text='Columnas')
        canvas = tk.Canvas(map_tab, highlightthickness=0, bg='#f6f4f7')
        scrollbar = ttk.Scrollbar(map_tab, orient='vertical', command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side='right', fill='y')
        canvas.pack(side='left', fill='both', expand=True)
        self.map_frame = ttk.Frame(canvas, padding=14)
        inner = canvas.create_window((0, 0), window=self.map_frame, anchor='nw')
        self.map_frame.bind('<Configure>', lambda e: canvas.configure(scrollregion=canvas.bbox('all')))
        canvas.bind('<Configure>', lambda e: canvas.itemconfigure(inner, width=e.width))
        self.preview_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.preview_tab, text='Revisión')
        self.preview = tk.Text(self.preview_tab, wrap='word', font=('Segoe UI', 11), padx=14, pady=14, state='disabled')
        preview_scroll = ttk.Scrollbar(self.preview_tab, command=self.preview.yview)
        self.preview.configure(yscrollcommand=preview_scroll.set)
        preview_scroll.pack(side='right', fill='y')
        self.preview.pack(fill='both', expand=True)
        footer = ttk.Frame(self, padding=(22, 14))
        footer.pack(fill='x')
        ttk.Button(footer, text='3. Revisar e importar', style='Accent.TButton', command=self.review).pack(side='right')
        ttk.Button(footer, text='Cerrar', command=self.close).pack(side='right', padx=10)

    def choose_file(self):
        path = filedialog.askopenfilename(parent=self, title='Seleccionar Excel de empleados', filetypes=[('Excel', '*.xlsx *.xlsm')])
        if not path:
            return
        try:
            source = ExcelSource(path)
        except ImportError:
            messagebox.showerror('Falta el lector de Excel', 'Ejecuta INSTALAR.bat una vez para instalar openpyxl.', parent=self)
            return
        except Exception as error:
            messagebox.showerror('No se pudo leer el archivo', str(error), parent=self)
            return
        if self.source:
            self.source.close()
        self.source = source
        self.filename.set(source.path.name)
        self.sheets.configure(values=source.sheets)
        self.sheet.set(source.sheets[0])
        self.header_row.set('1')
        self.read_columns()

    def read_columns(self):
        self.loaded_config = None
        if not self.source:
            messagebox.showinfo('Selecciona un archivo', 'Primero selecciona el Excel.', parent=self)
            return
        try:
            headers = self.source.headers(self.sheet.get(), int(self.header_row.get()))
        except Exception as error:
            messagebox.showerror('Revisa la hoja y la fila', str(error), parent=self)
            return
        from openpyxl.utils import get_column_letter
        choices = ['(Sin importar)'] + [f'{get_column_letter(i + 1)} — {name or "Sin título"}' for i, name in enumerate(headers)]
        self.column_indexes = {choice: index for index, choice in enumerate(choices[1:])}
        mapping = suggest_mapping(headers)
        for child in self.map_frame.winfo_children():
            child.destroy()
        self.mapping_vars = {}
        for index, (key, label, _) in enumerate(FIELDS):
            ttk.Label(self.map_frame, text=label + (' *' if key in REQUIRED else '')).grid(row=index, column=0, sticky='w', padx=(0, 18), pady=5)
            var = tk.StringVar(value=choices[mapping[key] + 1] if key in mapping else choices[0])
            self.mapping_vars[key] = var
            ttk.Combobox(self.map_frame, textvariable=var, values=choices, state='readonly', width=48).grid(row=index, column=1, sticky='ew', pady=5)
        self.map_frame.columnconfigure(1, weight=1)
        self.loaded_config = (self.sheet.get(), int(self.header_row.get()))
        self.notebook.select(0)

    def show_review(self, text):
        self.preview.configure(state='normal')
        self.preview.delete('1.0', 'end')
        self.preview.insert('1.0', text)
        self.preview.configure(state='disabled')
        self.notebook.select(self.preview_tab)

    def review(self):
        try:
            config = (self.sheet.get(), int(self.header_row.get()))
            if not self.source or self.loaded_config != config:
                raise ValueError('Pulsa Leer columnas después de elegir la hoja y la fila de encabezados.')
            mapping = {k: self.column_indexes[v.get()] for k, v in self.mapping_vars.items() if v.get() in self.column_indexes}
            records, errors = self.source.records(*config, mapping, self.default_region.get())
        except Exception as error:
            messagebox.showerror('No se puede importar', str(error), parent=self)
            return
        missing = [LABELS[k] for k in KEYS if k not in mapping and k not in ('regional', 'numero')]
        used = set(mapping.values())
        ignored = [choice for choice, index in self.column_indexes.items() if index not in used]
        summary = [f'Hoja: {config[0]}', f'Filas válidas: {len(records)}', f'Filas con errores: {len(errors)}',
                   'Regional: ' + ('según la columna del Excel' if 'regional' in mapping else self.default_region.get()),
                   ' · '.join(f'{region}: {sum(data.get("regional") == region for _, data in records)}' for region in REGIONAL_OPTIONS),
                   'Asignaciones: ' + str(sum(len(raw.get('_asignaciones', [])) for _,raw in records)),
                   'Nombres pendientes de separar: ' + str(sum(not raw.get('apellidos') or not raw.get('nombres') for _,raw in records)),
                   'Sin distribución de horas: ' + str(sum(not raw.get('_asignaciones') for _,raw in records)),
                   'Número de empleado: se conserva si está en Excel; si falta, se asigna al importar.',
                   'Campos sin asignar: ' + (', '.join(missing) or 'Ninguno'),
                   'Columnas que no se importarán: ' + (', '.join(ignored) or 'Ninguna'), '']
        if errors:
            summary += ['No se importará ninguna fila hasta corregir estos errores.', *errors]
        else:
            summary += ['VISTA PREVIA — primeros 5 empleados. Los campos vacíos se muestran como —.']
            for row_number, data in records[:5]:
                canonical = validate(data)
                summary += [f'\nFila {row_number}'] + [f'{label}: {display(k, canonical.get(k)) or ("Automático" if k == "numero" else "—")}' for k, label, _ in FIELDS]
                summary += ['  Asignación: ' + ' | '.join(row[k] for k in ASSIGNMENT_KEYS) for row in data.get('_asignaciones', [])]
        self.show_review('\n'.join(summary))
        if errors or not records:
            return
        # La revisión queda visible antes de la confirmación de escritura.
        if not messagebox.askyesno('Revisar vista previa',
                f'Se encontraron {len(records)} filas válidas.\nCampos sin asignar: {len(missing)}. Columnas omitidas: {len(ignored)}.\n\n¿Continuar con la importación?\nSi necesitas leer la vista previa completa, pulsa No.\nLos duplicados exactos se omiten; cualquier conflicto cancela todo el lote.', parent=self):
            return
        try:
            backups = self.app.db.path.parent / 'copias_importacion'
            backups.mkdir(exist_ok=True)
            backup = backups / ('antes_importar_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.db')
            self.app.db.backup(backup)
            added, skipped = self.app.db.import_rows(records)
        except Exception as error:
            self.show_review('\n'.join(summary) + '\n\nIMPORTACIÓN CANCELADA: ' + str(error))
            messagebox.showerror('Importación cancelada', f'No se guardó ningún registro del lote.\n\n{error}', parent=self)
            return
        self.app.clear_search()
        messagebox.showinfo('Importación completada', f'Empleados agregados: {added}\nDuplicados exactos omitidos: {skipped}\n\nCopia previa guardada en:\n{backup}', parent=self)
        self.close()

    def close(self):
        if self.source:
            self.source.close()
        self.destroy()
